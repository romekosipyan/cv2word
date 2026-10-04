"""Scoring helpers for the US-070 fidelity spike.

Character scores are whitespace-normalized multisets. Accents and digits stay.
Do not blend scores across layout classes.
"""

from __future__ import annotations

import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

WS_RE = re.compile(r"\s+")
# Formatting artifacts, not content: soft hyphen, word joiner, zero-width marks.
INVISIBLE_RE = re.compile(r"[\u00ad\u200b\u200c\u200d\u2060\ufeff]")
DASH_RE = re.compile(r"[\u2010\u2011\u2012\u2013\u2014\u2212]")
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def canonicalize(text: str) -> str:
    """Whitespace + hyphen unification. Accents and digits stay."""
    text = unicodedata.normalize("NFKC", text or "")
    text = INVISIBLE_RE.sub("", text)
    text = DASH_RE.sub("-", text)
    text = "".join(" " if unicodedata.category(ch) == "Zs" else ch for ch in text)
    return text


def normalize_chars(text: str) -> str:
    """Drop whitespace only after canonicalization. Keep case, accents, digits."""
    return WS_RE.sub("", canonicalize(text))


def char_multiset(text: str) -> Counter[str]:
    return Counter(normalize_chars(text))


def character_recall(expected: str, actual: str) -> float:
    exp = char_multiset(expected)
    act = char_multiset(actual)
    if not exp:
        return 1.0
    matched = sum(min(exp[ch], act[ch]) for ch in exp)
    return matched / sum(exp.values())


def character_precision(expected: str, actual: str) -> float:
    exp = char_multiset(expected)
    act = char_multiset(actual)
    if not act:
        return 0.0
    matched = sum(min(exp[ch], act[ch]) for ch in act)
    return matched / sum(act.values())


def count_occurrences(haystack: str, needle: str) -> int:
    if not needle:
        return 0
    return canonicalize(haystack).casefold().count(canonicalize(needle).casefold())


def phone_digits(value: str) -> str:
    return "".join(ch for ch in canonicalize(value) if ch.isdigit())


def fact_present(haystack: str, fact: str, kind: str | None = None) -> bool:
    if count_occurrences(haystack, fact) > 0:
        return True
    if kind == "phone":
        digits = phone_digits(fact)
        return bool(digits) and digits in phone_digits(haystack)
    return False


def duplicate_flags(expected: str, actual: str, facts: list[str]) -> list[str]:
    flags: list[str] = []
    exp_n = len(normalize_chars(expected))
    act_n = len(normalize_chars(actual))
    if exp_n and act_n > exp_n * 1.45:
        flags.append("output_much_longer_than_source")
    for fact in facts:
        if not fact:
            continue
        if count_occurrences(actual, fact) > max(1, count_occurrences(expected, fact)):
            flags.append(f"duplicated:{fact}")
    return flags


def _xml_text(xml_bytes: bytes) -> str:
    root = etree.fromstring(xml_bytes)
    runs = [node.text or "" for node in root.iter(f"{{{W_NS}}}t")]
    return "\n".join(run for run in runs if run)


def extract_docx_text(path: str) -> str:
    """Read every w:t run, including text boxes python-docx paragraphs miss."""
    parts: list[str] = []
    with ZipFile(path) as zf:
        for name in zf.namelist():
            if name.startswith("word/") and name.endswith(".xml") and "rels" not in name:
                if any(key in name for key in ("document", "header", "footer", "footnotes", "endnotes")):
                    parts.append(_xml_text(zf.read(name)))
    if parts:
        return "\n".join(parts)
    doc = Document(path)
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


def count_docx_images(path: str) -> int:
    doc = Document(path)
    return len(doc.element.findall(".//" + qn("a:blip")))


def docx_has_native_body_text(path: str, min_chars: int = 80) -> bool:
    return len(normalize_chars(extract_docx_text(path))) >= min_chars


def docx_is_image_only_pages(path: str, min_native_chars: int = 80) -> bool:
    native = len(normalize_chars(extract_docx_text(path)))
    images = count_docx_images(path)
    return images > 0 and native < min_native_chars


NSMAP = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def docx_package_ok(path: str) -> bool:
    try:
        with ZipFile(path) as zf:
            names = set(zf.namelist())
            if "word/document.xml" not in names:
                return False
            xml = zf.read("word/document.xml")
        etree.fromstring(xml)
        return True
    except Exception:
        return False


def _loose_alnum(text: str) -> str:
    return "".join(ch for ch in canonicalize(text).casefold() if ch.isalnum())


def column_interleave(actual: str, left_only: list[str], right_only: list[str]) -> dict:
    events: list[tuple[int, str]] = []
    haystack = _loose_alnum(actual)
    for token in left_only:
        needle = _loose_alnum(token)
        idx = haystack.find(needle) if needle else -1
        if idx >= 0:
            events.append((idx, "L"))
    for token in right_only:
        needle = _loose_alnum(token)
        idx = haystack.find(needle) if needle else -1
        if idx >= 0:
            events.append((idx, "R"))
    events.sort()
    sequence = "".join(side for _, side in events)
    switches = sum(1 for a, b in zip(sequence, sequence[1:]) if a != b)
    interleaved = switches > 1 and ("LRL" in sequence or "RLR" in sequence)
    return {
        "sequence": sequence,
        "switches": switches,
        "interleaved": bool(interleaved),
        "left_found": sum(1 for _, side in events if side == "L"),
        "right_found": sum(1 for _, side in events if side == "R"),
    }


@dataclass
class CriticalFacts:
    name: str
    email: str
    phone: str
    dates: list[str] = field(default_factory=list)
    employers: list[str] = field(default_factory=list)

    def all_items(self) -> list[tuple[str, str]]:
        items = [("name", self.name), ("email", self.email), ("phone", self.phone)]
        items.extend(("date", value) for value in self.dates)
        items.extend(("employer", value) for value in self.employers)
        return items

    def check(self, actual: str, source: str | None = None) -> dict:
        results = {}
        missing = []
        checked = 0
        for kind, value in self.all_items():
            if source is not None and not fact_present(source, value, kind):
                continue
            checked += 1
            ok = fact_present(actual, value, kind)
            results[f"{kind}:{value}"] = ok
            if not ok:
                missing.append(kind)
        return {
            "all_present": checked > 0 and not missing,
            "missing_kinds": sorted(set(missing)),
            "checked": checked,
        }

    def values(self) -> list[str]:
        return [value for _, value in self.all_items()]
