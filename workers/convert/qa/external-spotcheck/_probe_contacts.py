from pathlib import Path
import re
import pymupdf
from docx import Document

root = Path(__file__).resolve().parent
EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")


def docx_text(path: Path) -> str:
    d = Document(str(path))
    parts = []
    for p in d.paragraphs:
        if p.text:
            parts.append(p.text)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


for sid in ["ext-02", "ext-04", "ext-05", "ext-07", "ext-08", "ext-10"]:
    pdf = next((root / "selected").glob(f"{sid}-*.pdf"))
    docx = next((root / "docx").glob(f"{sid}-*.docx"))
    pt = "\n".join((page.get_text("text") or "") for page in pymupdf.open(pdf))
    dt = docx_text(docx)
    emails = EMAIL_RE.findall(pt)
    phones = PHONE_RE.findall(pt)
    print("===", sid, pdf.name)
    print(" pdf emails", len(emails), "domains", [e.split("@")[-1] for e in emails[:3]])
    print(" pdf phones", len(phones), "digit_lens", [len(re.sub(r"\D", "", p)) for p in phones[:3]])
    for e in emails[:2]:
        local, _, domain = e.partition("@")
        print(
            "  domain in docx?",
            domain.lower() in dt.lower(),
            "local in docx?",
            local.lower() in dt.lower(),
        )
    for p in phones[:2]:
        dig = re.sub(r"\D", "", p)
        print("  phone digits in docx?", dig in re.sub(r"\D", "", dt), "len", len(dig))
    print("  docx has @?", "@" in dt, "docx emails", [e.split("@")[-1] for e in EMAIL_RE.findall(dt)[:2]])
    # show whether contact may be glyph/icon separated
    contactish = [ln for ln in pt.splitlines() if "@" in ln or re.search(r"\d{3}", ln)][:6]
    print("  pdf contactish line lens", [len(x.strip()) for x in contactish])
