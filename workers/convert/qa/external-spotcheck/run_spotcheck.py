"""External resume PDF fidelity spot-check (QA).

Uses the same pdf2docx convert kwargs as production client/worker:
parse_stream_table=False, float_image_ignorable_gap=2.0,
line_separate_threshold=3.0, connected_border_tolerance=0.2.

Does not dump resume body text into reports.
"""
from __future__ import annotations

import json
import re
import shutil
import traceback
from datetime import datetime, timezone
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml.ns import qn
from pdf2docx import Converter

ROOT = Path(__file__).resolve().parent
RAW = ROOT / "pdfs"
PDFS = ROOT / "selected"
DOCX = ROOT / "docx"
SHOTS = ROOT / "screenshots"
MANIFEST_PATH = ROOT / "manifest.json"
EVAL_PATH = ROOT / "eval.json"

# Curated 10 from rights-safe open-source template demos (internet samples).
# Mix: simple ~4, two-column/sidebar ~4, complex/graphic ~2–3.
CURATED = [
    {
        "id": "ext-01",
        "source_file": "latexcv-classic.pdf",
        "file": "ext-01-latexcv-classic.pdf",
        "source_url": "https://raw.githubusercontent.com/jankapunkt/latexcv/master/classic/main.pdf",
        "layout_class": "simple",
        "notes": "latexcv classic — mostly single-column text demo (MIT)",
    },
    {
        "id": "ext-02",
        "source_file": "latexcv-minimalistic.pdf",
        "file": "ext-02-latexcv-minimalistic.pdf",
        "source_url": "https://raw.githubusercontent.com/jankapunkt/latexcv/master/minimalistic/main.pdf",
        "layout_class": "simple",
        "notes": "latexcv minimalistic — clean single-column (MIT)",
    },
    {
        "id": "ext-03",
        "source_file": "latexcv-rows.pdf",
        "file": "ext-03-latexcv-rows.pdf",
        "source_url": "https://raw.githubusercontent.com/jankapunkt/latexcv/master/rows/main.pdf",
        "layout_class": "complex_graphic",
        "notes": "latexcv rows — colored bands, photo, vertical labels (MIT); visually rich",
    },
    {
        "id": "ext-04",
        "source_file": "ice1000-resume.pdf",
        "file": "ext-04-ice1000-resume.pdf",
        "source_url": "https://raw.githubusercontent.com/ice1000/resume/master/resume.pdf",
        "layout_class": "simple",
        "notes": "ice1000 open-source resume template demo (mostly text)",
    },
    {
        "id": "ext-05",
        "source_file": "deedy-two-column.pdf",
        "file": "ext-05-deedy-two-column.pdf",
        "source_url": "https://raw.githubusercontent.com/deedydas/Deedy-Resume/master/OpenFonts/deedy_resume-openfont.pdf",
        "layout_class": "two_column",
        "notes": "Deedy Resume OpenFonts — asymmetric two-column (demo template)",
    },
    {
        "id": "ext-06",
        "source_file": "latexcv-two-column.pdf",
        "file": "ext-06-latexcv-two-column.pdf",
        "source_url": "https://raw.githubusercontent.com/jankapunkt/latexcv/master/two_column/main.pdf",
        "layout_class": "two_column",
        "notes": "latexcv two_column demo (MIT)",
    },
    {
        "id": "ext-07",
        "source_file": "latexcv-sidebar.pdf",
        "file": "ext-07-latexcv-sidebar.pdf",
        "source_url": "https://raw.githubusercontent.com/jankapunkt/latexcv/master/sidebar/main.pdf",
        "layout_class": "two_column",
        "notes": "latexcv sidebar — colored sidebar layout (MIT)",
    },
    {
        "id": "ext-08",
        "source_file": "awesome-cv-resume.pdf",
        "file": "ext-08-awesome-cv-resume.pdf",
        "source_url": "https://raw.githubusercontent.com/posquit0/Awesome-CV/master/examples/resume.pdf",
        "layout_class": "simple",
        "notes": "Awesome-CV example — single-column with icon header (LPPL)",
    },
    {
        "id": "ext-09",
        "source_file": "latexcv-modern.pdf",
        "file": "ext-09-latexcv-modern.pdf",
        "source_url": "https://raw.githubusercontent.com/jankapunkt/latexcv/master/modern/main.pdf",
        "layout_class": "complex_graphic",
        "notes": "latexcv modern — photo, QR, colored header (MIT); visually rich",
    },
    {
        "id": "ext-10",
        "source_file": "latexcv-infographics.pdf",
        "file": "ext-10-latexcv-infographics.pdf",
        "source_url": "https://raw.githubusercontent.com/jankapunkt/latexcv/master/infographics/main.pdf",
        "layout_class": "complex_graphic",
        "notes": "latexcv infographics — charts/graphics heavy (MIT); visually rich",
    },
]


EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")
DATE_RE = re.compile(
    r"\b(?:19|20)\d{2}\b|\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\b",
    re.I,
)


def page_size_label(w: float, h: float) -> str:
    # points
    def near(a: float, b: float, tol: float = 12) -> bool:
        return abs(a - b) <= tol

    if near(w, 595, 20) and near(h, 842, 20):
        return "A4"
    if near(w, 612, 20) and near(h, 792, 20):
        return "Letter"
    if near(w, 842, 20) and near(h, 595, 20):
        return "A4-landscape"
    return f"custom_{w:.0f}x{h:.0f}"


def inspect_pdf(path: Path) -> dict:
    doc = pymupdf.open(path)
    pages = doc.page_count
    w = doc[0].rect.width
    h = doc[0].rect.height
    total_chars = 0
    total_images = 0
    total_drawings = 0
    fonts: set[str] = set()
    for page in doc:
        text = page.get_text("text") or ""
        total_chars += len(text.strip())
        blocks = page.get_text("dict")["blocks"]
        total_images += sum(1 for b in blocks if b.get("type") == 1)
        total_drawings += len(page.get_drawings())
        for b in blocks:
            if b.get("type") != 0:
                continue
            for line in b.get("lines", []):
                for span in line.get("spans", []):
                    if span.get("font"):
                        fonts.add(str(span["font"]))
    # render first page preview
    pix = doc[0].get_pixmap(matrix=pymupdf.Matrix(1.2, 1.2), alpha=False)
    preview = SHOTS / f"{path.stem}-pdf-p1.png"
    pix.save(str(preview))
    doc.close()
    return {
        "pages": pages,
        "page_size": page_size_label(w, h),
        "width_pt": round(w, 1),
        "height_pt": round(h, 1),
        "pdf_chars": total_chars,
        "images": total_images,
        "drawings": total_drawings,
        "font_count": len(fonts),
        "preview": str(preview.relative_to(ROOT)),
    }


def convert_pdf(pdf_path: Path, docx_path: Path) -> dict:
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    if docx_path.exists():
        docx_path.unlink()
    started = datetime.now(timezone.utc)
    try:
        cv = Converter(str(pdf_path))
        try:
            cv.convert(
                str(docx_path),
                parse_stream_table=False,
                float_image_ignorable_gap=2.0,
                line_separate_threshold=3.0,
                connected_border_tolerance=0.2,
            )
        finally:
            cv.close()
        ok = docx_path.is_file() and docx_path.stat().st_size > 0
        return {
            "ok": ok,
            "error": None if ok else "docx_missing_or_empty",
            "bytes": docx_path.stat().st_size if ok else 0,
            "elapsed_s": (datetime.now(timezone.utc) - started).total_seconds(),
        }
    except Exception as exc:  # noqa: BLE001 — spot-check harness
        return {
            "ok": False,
            "error": f"{type(exc).__name__}: {exc}",
            "bytes": 0,
            "elapsed_s": (datetime.now(timezone.utc) - started).total_seconds(),
            "traceback": traceback.format_exc(limit=3),
        }


def normalize(s: str) -> str:
    return re.sub(r"\s+", "", s or "").lower()


def docx_text(path: Path) -> str:
    d = Document(str(path))
    parts: list[str] = []
    for p in d.paragraphs:
        if p.text:
            parts.append(p.text)
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def pdf_text(path: Path) -> str:
    doc = pymupdf.open(path)
    parts = [(page.get_text("text") or "") for page in doc]
    doc.close()
    return "\n".join(parts)


def fact_presence(pdf_t: str, docx_t: str) -> dict:
    """Check critical-fact *presence* without storing PII values."""
    emails = EMAIL_RE.findall(pdf_t)
    phones = PHONE_RE.findall(pdf_t)
    dates = DATE_RE.findall(pdf_t)
    # names/employers: use first non-empty line tokens longer than 3 chars as weak signals
    lines = [ln.strip() for ln in pdf_t.splitlines() if ln.strip()]
    name_line = lines[0] if lines else ""
    # employers: look for common section content lines under Experience-ish
    employer_candidates = [
        ln
        for ln in lines
        if 4 <= len(ln) <= 60
        and not EMAIL_RE.search(ln)
        and not PHONE_RE.search(ln)
        and not ln.lower().startswith(("experience", "education", "skills", "summary"))
    ][:8]

    def present(values: list[str]) -> tuple[int, int]:
        if not values:
            return 0, 0
        hit = sum(1 for v in values if normalize(v) and normalize(v) in normalize(docx_t))
        return hit, len(values)

    email_h, email_n = present(emails[:3])
    phone_h, phone_n = present([re.sub(r"\s+", "", p) for p in phones[:3]])
    # phones may differ by separators — also check digit-only
    if phone_n and phone_h < phone_n:
        dig_pdf = [re.sub(r"\D", "", p) for p in phones[:3]]
        dig_docx = re.sub(r"\D", "", docx_t)
        phone_h = sum(1 for d in dig_pdf if d and d in dig_docx)
    date_h, date_n = present(list(dict.fromkeys(dates))[:12])
    name_ok = bool(name_line) and normalize(name_line[:40])[:12] in normalize(docx_t)
    emp_h, emp_n = present(employer_candidates[:5])

    return {
        "email_preserved": None if email_n == 0 else email_h == email_n,
        "email_hits": f"{email_h}/{email_n}",
        "phone_preserved": None if phone_n == 0 else phone_h == phone_n,
        "phone_hits": f"{phone_h}/{phone_n}",
        "dates_mostly_preserved": None if date_n == 0 else (date_h / date_n) >= 0.8,
        "date_hits": f"{date_h}/{date_n}",
        "name_line_preserved": name_ok,
        "employer_like_hits": f"{emp_h}/{emp_n}",
        "employers_mostly_preserved": None if emp_n == 0 else (emp_h / emp_n) >= 0.6,
    }


def char_stats(pdf_t: str, docx_t: str) -> dict:
    a = normalize(pdf_t)
    b = normalize(docx_t)
    if not a:
        return {"recall": None, "precision": None, "pdf_norm_chars": 0, "docx_norm_chars": len(b)}
    # bag-of-char recall/precision (order-insensitive) for qualitative completeness
    from collections import Counter

    ca, cb = Counter(a), Counter(b)
    inter = sum((ca & cb).values())
    recall = inter / sum(ca.values())
    precision = inter / sum(cb.values()) if cb else 0.0
    return {
        "recall": round(recall, 4),
        "precision": round(precision, 4),
        "pdf_norm_chars": sum(ca.values()),
        "docx_norm_chars": sum(cb.values()),
    }


def reading_order_heuristic(pdf_t: str, docx_t: str, layout_class: str) -> dict:
    """Very light interleave heuristic: section-header order vs appearance in DOCX."""
    headers = []
    for ln in pdf_t.splitlines():
        s = ln.strip()
        if not s or len(s) > 40:
            continue
        if s.isupper() or s.istitle():
            if s.lower() in {
                "experience",
                "education",
                "skills",
                "projects",
                "summary",
                "work experience",
                "professional experience",
                "publications",
                "awards",
                "languages",
                "interests",
                "contact",
                "profile",
                "employment",
            } or (s.isupper() and len(s) >= 5):
                headers.append(s)
    # unique preserve order
    seen = set()
    ordered = []
    for h in headers:
        k = h.lower()
        if k not in seen:
            seen.add(k)
            ordered.append(h)
    ordered = ordered[:8]
    positions = []
    nd = normalize(docx_t)
    for h in ordered:
        p = nd.find(normalize(h))
        positions.append(p)
    increasing = all(
        positions[i] <= positions[i + 1]
        for i in range(len(positions) - 1)
        if positions[i] >= 0 and positions[i + 1] >= 0
    )
    missing = sum(1 for p in positions if p < 0)
    suspect_interleave = layout_class in {"two_column", "complex_graphic"} and (
        not increasing or missing >= 2
    )
    return {
        "section_headers_checked": len(ordered),
        "section_headers_found_in_docx": sum(1 for p in positions if p >= 0),
        "header_order_nondecreasing": increasing if ordered else None,
        "suspect_column_interleave": suspect_interleave,
    }


def evaluate_docx(pdf_path: Path, docx_path: Path, layout_class: str) -> dict:
    if not docx_path.is_file():
        return {"editable_native_text": False, "openable": False}
    d = Document(str(docx_path))
    blips = d.element.body.findall(".//" + qn("a:blip"))
    para_text_len = sum(len(p.text or "") for p in d.paragraphs)
    table_text_len = 0
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                table_text_len += len(cell.text or "")
    body_len = para_text_len + table_text_len
    pdf_t = pdf_text(pdf_path)
    docx_t = docx_text(docx_path)
    stats = char_stats(pdf_t, docx_t)
    facts = fact_presence(pdf_t, docx_t)
    order = reading_order_heuristic(pdf_t, docx_t, layout_class)
    # link check
    pdf_links = 0
    doc = pymupdf.open(pdf_path)
    for page in doc:
        pdf_links += len(page.get_links() or [])
    doc.close()
    # DOCX hyperlinks
    hyper = d.element.body.findall(".//" + qn("w:hyperlink"))
    # visual flags (heuristic from structure)
    visual = {
        "many_tables": len(d.tables) >= 8,
        "image_count_docx": len(blips),
        "image_count_pdf": None,  # filled by caller
        "likely_spacing_or_structure_heavy": len(d.tables) >= 5 or len(blips) == 0,
    }
    return {
        "openable": True,
        "editable_native_text": body_len >= 80,
        "docx_paras": len(d.paragraphs),
        "docx_tables": len(d.tables),
        "docx_images": len(blips),
        "body_text_chars": body_len,
        "char_completeness": stats,
        "critical_facts": facts,
        "reading_order": order,
        "pdf_link_count": pdf_links,
        "docx_hyperlink_count": len(hyper),
        "links_note": (
            "no_pdf_links"
            if pdf_links == 0
            else (
                "links_present_in_docx"
                if len(hyper) > 0
                else "pdf_had_links_docx_hyperlinks_missing_or_plain"
            )
        ),
        "visual_structure": visual,
    }


def main() -> None:
    PDFS.mkdir(parents=True, exist_ok=True)
    DOCX.mkdir(parents=True, exist_ok=True)
    SHOTS.mkdir(parents=True, exist_ok=True)

    manifest = []
    evals = []
    for item in CURATED:
        src = RAW / item["source_file"]
        if not src.is_file():
            raise SystemExit(f"missing source {src}")
        dest = PDFS / item["file"]
        shutil.copy2(src, dest)
        meta = inspect_pdf(dest)
        entry = {
            **item,
            "origin": "internet_sample",
            **meta,
        }
        manifest.append(entry)

        out = DOCX / f"{Path(item['file']).stem}.docx"
        conv = convert_pdf(dest, out)
        ev = {
            "id": item["id"],
            "file": item["file"],
            "layout_class": item["layout_class"],
            "conversion": conv,
        }
        if conv["ok"]:
            detail = evaluate_docx(dest, out, item["layout_class"])
            detail["visual_structure"]["image_count_pdf"] = meta["images"]
            # missing icons/images heuristic
            detail["visual_issues"] = {
                "images_dropped": meta["images"] > 0 and detail["docx_images"] == 0,
                "images_reduced": meta["images"] > 0
                and detail["docx_images"] < meta["images"],
                "complex_drawings_in_pdf": meta["drawings"] >= 40,
                "many_bogus_tables_risk": detail["docx_tables"] >= 10,
            }
            ev["evaluation"] = detail
        else:
            ev["evaluation"] = None
        evals.append(ev)
        print(
            item["id"],
            item["layout_class"],
            "convert",
            conv["ok"],
            conv.get("error"),
            "pages",
            meta["pages"],
            meta["page_size"],
        )

    MANIFEST_PATH.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    EVAL_PATH.write_text(
        json.dumps(
            {
                "generated_at": datetime.now(timezone.utc).isoformat(),
                "converter_path": "local Python pdf2docx==0.5.13 via workers/convert/.venv "
                "(mirrors apps/web pyodideConvert + engine/convert.py kwargs)",
                "pymupdf": "1.28.2 (worker pin; client Pyodide uses 1.27.2.3 wheel)",
                "results": evals,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print("wrote", MANIFEST_PATH, EVAL_PATH)


if __name__ == "__main__":
    main()
