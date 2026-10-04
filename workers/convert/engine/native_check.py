"""Native-text hooks for the convert path (US-010 / US-011).

US-011 `engine.validate` owns package/macro/metadata/page completeness.
This module counts editable `w:t` runs so success never means page images
without editable text. Never returns extracted resume text to JSON callers.
"""

from __future__ import annotations

import re
from pathlib import Path
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.oxml.ns import qn

WS_RE = re.compile(r"\s+")
MIN_NATIVE_CHARS = 80
W_T = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"


def _normalize_chars(text: str) -> str:
    return WS_RE.sub("", text or "")


def docx_package_has_document(path: Path) -> bool:
    try:
        with ZipFile(path) as zf:
            return "word/document.xml" in zf.namelist()
    except BadZipFile:
        return False


def count_docx_images(path: Path) -> int:
    doc = Document(str(path))
    return len(doc.element.findall(".//" + qn("a:blip")))


def native_body_char_count(path: Path) -> int:
    """Count non-whitespace characters from w:t runs (incl. table cells).

    pdf2docx often places body text in tables; paragraph-only counts miss that.
    Returns a length only — never the text itself.
    """
    total = 0
    with ZipFile(path) as zf:
        for name in zf.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            if "rels" in name:
                continue
            if not any(
                key in name
                for key in ("document", "header", "footer", "footnotes", "endnotes")
            ):
                continue
            root = ElementTree.fromstring(zf.read(name))
            for node in root.iter(W_T):
                total += len(_normalize_chars(node.text or ""))
    if total > 0:
        return total
    # Fallback for unusual packages python-docx can still open.
    doc = Document(str(path))
    joined = "\n".join(p.text for p in doc.paragraphs)
    return len(_normalize_chars(joined))


def native_editable_ok(path: Path, *, min_chars: int = MIN_NATIVE_CHARS) -> dict:
    """Return a JSON-safe check result without document text."""
    if not path.is_file() or path.stat().st_size <= 0:
        return {
            "ok": False,
            "error": "output_invalid",
            "reason": "missing_or_empty",
            "native_chars": 0,
            "image_count": 0,
        }
    if not docx_package_has_document(path):
        return {
            "ok": False,
            "error": "output_invalid",
            "reason": "package_invalid",
            "native_chars": 0,
            "image_count": 0,
        }
    try:
        native = native_body_char_count(path)
        images = count_docx_images(path)
    except Exception:  # noqa: BLE001 — treat unreadable DOCX as invalid output
        return {
            "ok": False,
            "error": "output_invalid",
            "reason": "unreadable",
            "native_chars": 0,
            "image_count": 0,
        }

    if native < min_chars:
        return {
            "ok": False,
            "error": "output_invalid",
            "reason": "image_only_or_empty",
            "native_chars": native,
            "image_count": images,
        }
    return {
        "ok": True,
        "native_chars": native,
        "image_count": images,
    }
